from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


EXPORT_DIR = Path("exports")


def format_duration(duration):
    """
    Convert timedelta into HH:MM format.
    """

    total_seconds = int(duration.total_seconds())

    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60

    return f"{hours:02d}:{minutes:02d}"


def calculate_break_duration(breaks):

    total = timedelta()

    for item in breaks:

        if item["end_time"]:

            start = datetime.fromisoformat(
                item["start_time"]
            )

            end = datetime.fromisoformat(
                item["end_time"]
            )

            total += end - start

    return total


def export_month(database, year, month):

    EXPORT_DIR.mkdir(exist_ok=True)

    sessions = database.get_sessions_for_month(
        year,
        month
    )

    if not sessions:
        return None

    workbook = Workbook()
    worksheet = workbook.active

    worksheet.title = (
        f"{year}-{month:02d}"
    )

    # ---------------------------------------------
    # TITLE
    # ---------------------------------------------

    worksheet["A1"] = (
        f"Working Hours Report - {year}-{month:02d}"
    )

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    # ---------------------------------------------
    # HEADERS
    # ---------------------------------------------

    headers = [
        "Date",
        "Start",
        "End",
        "Break",
        "Net Working Hours",
        "Status / Notes"
    ]

    header_row = 3

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="4F81BD"
    )

    for column, header in enumerate(headers, 1):

        cell = worksheet.cell(
            row=header_row,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center"
        )

    # ---------------------------------------------
    # DATA
    # ---------------------------------------------

    row = 4
    monthly_total = timedelta()

    for session in sessions:

        start = datetime.fromisoformat(
            session["start_time"]
        )

        end = datetime.fromisoformat(
            session["end_time"]
        )

        gross_duration = end - start

        breaks = database.get_breaks_for_session(
            session["id"]
        )

        break_duration = calculate_break_duration(
            breaks
        )

        net_duration = (
            gross_duration - break_duration
        )

        monthly_total += net_duration

        notes = session["notes"] or ""

        if session["end_time_edited"]:

            if notes:
                notes += " | "

            notes += "End time manually edited."

        worksheet.cell(
            row=row,
            column=1,
            value=start.strftime("%d.%m.%Y")
        )

        worksheet.cell(
            row=row,
            column=2,
            value=start.strftime("%H:%M")
        )

        worksheet.cell(
            row=row,
            column=3,
            value=end.strftime("%H:%M")
        )

        worksheet.cell(
            row=row,
            column=4,
            value=format_duration(
                break_duration
            )
        )

        worksheet.cell(
            row=row,
            column=5,
            value=format_duration(
                net_duration
            )
        )

        worksheet.cell(
            row=row,
            column=6,
            value=notes
        )

        row += 1

    # ---------------------------------------------
    # MONTHLY TOTAL
    # ---------------------------------------------

    row += 1

    worksheet.cell(
        row=row,
        column=4,
        value="MONTHLY TOTAL"
    )

    worksheet.cell(
        row=row,
        column=5,
        value=format_duration(
            monthly_total
        )
    )

    worksheet.cell(
        row=row,
        column=4
    ).font = Font(bold=True)

    worksheet.cell(
        row=row,
        column=5
    ).font = Font(bold=True)

    # ---------------------------------------------
    # COLUMN WIDTHS
    # ---------------------------------------------

    widths = {
        1: 15,
        2: 12,
        3: 12,
        4: 12,
        5: 20,
        6: 55
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    # ---------------------------------------------
    # SAVE
    # ---------------------------------------------

    filename = (
        EXPORT_DIR /
        f"work_hours_{year}-{month:02d}.xlsx"
    )

    workbook.save(filename)

    return filename